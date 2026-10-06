"""
SYNTHETIC TEST FIXTURE — generates small workbooks with the same LAYOUT as the HVPNL daily log sheets so the
pipeline and API can be tested in CI without the private files. These numbers are not electrical measurements,
are never shipped with the product and are never shown in the UI.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import numpy as np
from openpyxl import Workbook

HEADERS = [  # (group, sub, leaf)
    ("VOLTAGE", "220 kV Bus-I", None), ("VOLTAGE", "220 kV Bus-II", None), ("VOLTAGE", "11 kV T-I", None),
    ("220 kV FEEDERS LOAD", "220KV Test-Alpha CKT - I", None), ("220 kV FEEDERS LOAD", "220KV Test-Alpha CKT - II", None),
    ("220 kV FEEDERS LOAD", "LOAD IN MVA", "220/66KV 160MVA T/F T-1"),
    ("11 kV FEEDERS LOAD", "25/31.5 MVA T/F T- 2 I/C -I", None),
    ("11 kV FEEDERS LOAD", "Feeder-A", None), ("11 kV FEEDERS LOAD", "Feeder-B", None), ("11 kV FEEDERS LOAD", "Feeder-C", None),
    ("11 kV FEEDERS LOAD", "Feeder-D", None),
    ("TRANSFORMER TEMPERATURE", "160 MVA T/F T- 1", "O"), ("TRANSFORMER TEMPERATURE", "160 MVA T/F T- 1", "HV W"),
    ("Frequency", None, None), ("Weather", None, None),
]
DAYS = ["02.02.2026", "10.02.2026", "27.02.2026"]
PLANTED = {"voltage_zero": ("02.02.2026", 10), "feeders_zero": ("10.02.2026", 8)}


def _sheet(ws, day_idx: int, rng: np.random.Generator, title: str, filled: bool = True):
    ws["A1"] = "HARYANA VIDYUT PRASARAN NIGAM LIMITED (SYNTHETIC TEST FIXTURE)"
    ws["A3"] = f"Daily Log Sheet of {title} — synthetic fixture"
    ws["A4"] = "Time "
    col = 2
    groups: dict[str, list[int]] = {}
    for g, sub, leaf in HEADERS:
        ws.cell(4, col, g)
        groups.setdefault(g, []).append(col)
        if sub:
            ws.cell(5, col, sub)
        if leaf:
            ws.cell(6, col, leaf)
        col += 1
    for g, cols in groups.items():  # merged group headers like the real sheets
        if len(cols) > 1:
            ws.merge_cells(start_row=4, start_column=cols[0], end_row=4, end_column=cols[-1])
    for h in range(1, 25):
        r = 6 + h
        ws.cell(r, 1, dt.time(h % 24, 0) if h < 24 else dt.timedelta(days=1))
        if not filled:
            continue
        load = 0.6 + 0.4 * np.sin((h - 6) / 24 * 2 * np.pi) ** 2 + 0.03 * day_idx
        vals = [233 - 4 * load + rng.normal(0, .4), 232 - 4 * load + rng.normal(0, .4), 11.2 + rng.normal(0, .05),
                round(100 * load + rng.normal(0, 3)), round(100 * load + rng.normal(0, 3)), round(40 * load + rng.normal(0, 1), 2),
                round(200 * load + rng.normal(0, 5)), round(60 * load), round(50 * load), round(45 * load), round(40 * load),
                round(40 + 8 * load + rng.normal(0, .3)), round(42 + 9 * load + rng.normal(0, .3)), round(50 + rng.normal(0, .02), 2), "CLEAR"]
        day = DAYS[day_idx]
        if (day, h) == PLANTED["voltage_zero"]:
            vals[2] = 0
        if (day, h) == PLANTED["feeders_zero"]:
            vals[6] = vals[7] = vals[8] = vals[9] = 0
        if h == 5:
            vals[3] = "-"            # missing reading
        if h == 15 and day_idx == 1:
            vals[10] = "PTW"         # operator state inside a numeric column
        if h == 16:
            vals[4] = "ON"
        for i, v in enumerate(vals):
            ws.cell(r, 2 + i, v)
    ws.cell(32, 1, "FIRST SHIFT FROM ………….HRS")  # footer block must be ignored


def make_dataset(root: Path) -> Path:
    raw = root / "raw"
    raw.mkdir(parents=True, exist_ok=True)
    rng = np.random.default_rng(7)
    for i, d in enumerate(DAYS):
        wb = Workbook()
        ws = wb.active
        ws.title = "220 Test-A"
        _sheet(ws, i, rng, "220 kV Sub-Station, Test-A")
        _sheet(wb.create_sheet("66 Blank"), i, rng, "66 kV Sub-Station, Blank", filled=False)
        wb.save(raw / f"{d} -TS Faridabad.xlsx")
    return raw
